"""
Converts the data-analysis workbook (e.g. SCI_TDS_Analysis_Scenarios_2026-10-02.xlsx) into
data/analysis-scenarios.json, the file served by GET /api/analysis/scenarios.

Requires openpyxl (pip install openpyxl). The API itself only reads the generated JSON.

Usage:
    python scripts/extract_analysis_scenarios.py <workbook.xlsx> [-o data/analysis-scenarios.json]
"""
import argparse
import json
import re
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional

from openpyxl import load_workbook

DEFAULT_OUTPUT = Path(__file__).resolve().parent.parent / "data" / "analysis-scenarios.json"

SCENARIO_ID = re.compile(r"\bS\d{2}\b")
SHORT_TEST_CASE = re.compile(r"\b\d{5}\.\d{2}\b")
TEST_CASE_CELL = re.compile(r"^(TC_[\w.]+?\.\d{5}\.\d{2})\(([^)]*)\)(?:\s+–\s+(.+))?$")
CODED_VALUE = re.compile(r"^(\d+)\s+(.+)$")
NUMBER = re.compile(r"\d+(?:\.\d+)?")

# Canonical data-source types, matched against the start of the workbook text (e.g. "BLF (system variables)")
SOURCE_TYPES = [
    ("pdf", "PDF"),
    ("pcap", "pcap"),
    ("blf", "BLF"),
    ("write_log", "write log"),
    ("test_spec", "test spec"),
    ("telegram_xlsx", "telegram xlsx"),
    ("lua_dissector", "Lua dissector"),
]


def cell_text(value: Any) -> Optional[str]:
    """Cell value as stripped text; None for empty cells and '–' placeholders."""
    if value is None:
        return None
    text = str(value).strip()
    return None if text in ("", "–", "-") else text


def cell_lines(value: Any) -> List[str]:
    """Multi-line cell as a list of lines without bullet characters."""
    text = cell_text(value)
    if text is None:
        return []
    items = [line.strip().lstrip("•").strip() for line in text.splitlines()]
    return [item for item in items if item]


def source_type(label: str) -> str:
    for key, prefix in SOURCE_TYPES:
        if label.lower().startswith(prefix.lower()):
            return key
    raise ValueError(f"Unknown data source '{label}' - extend SOURCE_TYPES")


def scenario_ids(value: Any) -> List[str]:
    return SCENARIO_ID.findall(cell_text(value) or "")


def sheet_rows(ws, header_row: int, expected: List[str]) -> Iterator[Dict[str, Any]]:
    """Yields the rows below header_row as {header: value}; fails if an expected column is missing."""
    columns = {str(c.value).strip(): c.column for c in ws[header_row] if c.value is not None}
    missing = [name for name in expected if name not in columns]
    if missing:
        raise ValueError(f"Sheet '{ws.title}' is missing columns {missing}")
    for row in ws.iter_rows(min_row=header_row + 1):
        values = {name: row[columns[name] - 1].value for name in expected}
        if any(value is not None for value in values.values()):
            yield values


def find_row(ws, column: int, text: str) -> int:
    for row in ws.iter_rows():
        if cell_text(row[column - 1].value) == text:
            return row[0].row
    raise ValueError(f"Sheet '{ws.title}' has no row starting with '{text}'")


def read_overview(ws) -> Dict[str, Any]:
    def labelled(label: str) -> str:
        value = cell_text(ws.cell(find_row(ws, 1, label), 2).value)
        if value is None:
            raise ValueError(f"Overview field '{label}' is empty")
        return value

    test_cases = []
    for row in sheet_rows(ws, find_row(ws, 1, "#"), [
        "#", "Test case", "Report verdict", "CANoe time window [s]",
        "Where it failed", "Most probable root cause", "Scenarios",
    ]):
        if not isinstance(row["#"], int):
            break
        match = TEST_CASE_CELL.match(cell_text(row["Test case"]) or "")
        if not match:
            raise ValueError(f"Unexpected test case cell '{row['Test case']}'")
        window_start, window_end = (float(n) for n in NUMBER.findall(cell_text(row["CANoe time window [s]"])))
        test_cases.append({
            "number": row["#"],
            "test_case_id": match.group(1),
            "variant": match.group(2) or None,
            "title": match.group(3),
            "verdict": cell_text(row["Report verdict"]),
            "window_start_s": window_start,
            "window_end_s": window_end,
            "failure_point": cell_text(row["Where it failed"]),
            "root_cause": cell_text(row["Most probable root cause"]),
            "scenario_ids": scenario_ids(row["Scenarios"]),
        })

    return {
        "title": cell_text(ws["A1"].value),
        "test_run": {
            "name": labelled("Test run"),
            "overall_verdict": labelled("Overall verdict"),
            "sut": labelled("SUT"),
            "test_system": labelled("Test system"),
            "data_sources": labelled("Data sources"),
            "time_correlation": labelled("Time correlation"),
        },
        "test_cases": test_cases,
    }


def resolve_test_cases(text: Optional[str], known: Dict[str, str]) -> List[str]:
    """Maps short references such as '02288.01' to the full test case IDs of the overview."""
    resolved = []
    for short in SHORT_TEST_CASE.findall(text or ""):
        if short not in known:
            raise ValueError(f"Test case '{short}' is not listed in the overview")
        if known[short] not in resolved:
            resolved.append(known[short])
    return resolved


def read_scenarios(ws, known: Dict[str, str]) -> List[Dict[str, Any]]:
    scenarios = []
    for row in sheet_rows(ws, 1, [
        "ID", "Scenario / finding", "Category", "Test case(s)", "Data sources", "Severity", "Confidence",
        "Symptom (what the report shows)", "Evidence (timestamps = CANoe time; frame = pcap frame no.)",
        "Analysis / root cause", "Potential reasons", "Recommendation", "How it was found (method)",
    ]):
        scope = cell_text(row["Test case(s)"])
        labels = [label.strip() for label in (cell_text(row["Data sources"]) or "").split(",") if label.strip()]
        scenarios.append({
            "id": cell_text(row["ID"]),
            "title": cell_text(row["Scenario / finding"]),
            "category": cell_text(row["Category"]),
            "test_cases": resolve_test_cases(scope, known),
            "test_case_scope": scope,
            "applies_to_all_test_cases": bool(scope and scope.lower().startswith("all")),
            "data_sources": [{"type": source_type(label), "label": label} for label in labels],
            "severity": cell_text(row["Severity"]),
            "confidence": cell_text(row["Confidence"]),
            "symptom": cell_text(row["Symptom (what the report shows)"]),
            "evidence": cell_lines(row["Evidence (timestamps = CANoe time; frame = pcap frame no.)"]),
            "root_cause": cell_text(row["Analysis / root cause"]),
            "potential_reasons": cell_lines(row["Potential reasons"]),
            "recommendation": cell_text(row["Recommendation"]),
            "method": cell_text(row["How it was found (method)"]),
        })
    return scenarios


def read_timeline(ws, known: Dict[str, str]) -> List[Dict[str, Any]]:
    events = []
    for row in sheet_rows(ws, 1, [
        "CANoe time [s]", "Wall clock", "pcap frame", "Source", "Direction", "Event", "Test case / phase", "Comment",
    ]):
        phase = cell_text(row["Test case / phase"])
        test_cases = resolve_test_cases(phase, known)
        events.append({
            "canoe_time_s": float(row["CANoe time [s]"]),
            "wall_clock": cell_text(row["Wall clock"]),
            "pcap_frame": row["pcap frame"],
            "source": source_type(cell_text(row["Source"])),
            "direction": cell_text(row["Direction"]),
            "event": cell_text(row["Event"]),
            "phase": phase,
            "test_case_id": test_cases[0] if test_cases else None,
            "comment": cell_text(row["Comment"]),
            "scenario_ids": scenario_ids(row["Comment"]),
        })
    return events


def read_state_history(ws, section: str, known: Dict[str, str]) -> Dict[str, Any]:
    states = []
    coding_note = None
    for row in sheet_rows(ws, 1, [
        "CANoe time [s]", "pcap frame", "Belegungszustand (byte 43)", "Grundstellungsfähigkeit (byte 44)",
        "Axle count (bytes 45-46)", "Duration until next state [s]", "Test case", "Remark",
    ]):
        time_value = row["CANoe time [s]"]
        if not isinstance(time_value, (int, float)):
            coding_note = cell_text(time_value)  # explanatory note below the table
            continue
        occupancy = CODED_VALUE.match(cell_text(row["Belegungszustand (byte 43)"]))
        resettable = CODED_VALUE.match(cell_text(row["Grundstellungsfähigkeit (byte 44)"]))
        phase = cell_text(row["Test case"])
        test_cases = resolve_test_cases(phase, known)
        states.append({
            "canoe_time_s": float(time_value),
            "pcap_frame": row["pcap frame"],
            "occupancy_code": int(occupancy.group(1)),
            "occupancy": occupancy.group(2),
            "resettable_code": int(resettable.group(1)),
            "resettable": resettable.group(2),
            "axle_count": cell_text(row["Axle count (bytes 45-46)"]),
            "duration_note": cell_text(row["Duration until next state [s]"])
            if isinstance(row["Duration until next state [s]"], str) else None,
            "phase": phase,
            "test_case_id": test_cases[0] if test_cases else None,
            "remark": cell_text(row["Remark"]),
        })

    # Column F holds formulas (next time - this time) without cached values; the last row names the run end
    for state, following in zip(states, states[1:] + [None]):
        if following is not None:
            end = following["canoe_time_s"]
        else:
            end = float(NUMBER.findall(state["duration_note"] or "")[-1])
        state["until_s"] = end
        state["duration_s"] = round(end - state["canoe_time_s"], 3)

    return {"section": section, "coding_note": coding_note, "states": states}


def read_method(ws) -> Dict[str, Any]:
    steps = []
    for row in sheet_rows(ws, find_row(ws, 1, "Step"), ["Step", "Activity", "Details / commands"]):
        step = cell_text(row["Step"])
        if not (step and step.isdigit()):
            break
        steps.append({
            "step": int(step),
            "activity": cell_text(row["Activity"]),
            "details": cell_text(row["Details / commands"]),
        })

    questions = []
    for row in ws.iter_rows(min_row=find_row(ws, 1, "Open questions for the team") + 1):
        question_id = cell_text(row[1].value)
        if question_id and re.fullmatch(r"Q\d+", question_id):
            question = cell_text(row[2].value)
            questions.append({"id": question_id, "question": question, "scenario_ids": scenario_ids(question)})

    return {"steps": steps, "open_questions": questions}


def extract(workbook_path: Path) -> Dict[str, Any]:
    workbook = load_workbook(workbook_path, data_only=True)
    overview = read_overview(workbook["Overview"])
    known = {tc["test_case_id"][-8:]: tc["test_case_id"] for tc in overview["test_cases"]}
    section = re.search(r"GFM-A '([^']+)'", overview["test_run"]["sut"])

    return {
        "source_file": workbook_path.name,
        "title": overview["title"],
        "test_run": overview["test_run"],
        "test_cases": overview["test_cases"],
        "scenarios": read_scenarios(workbook["Analysis Scenarios"], known),
        "timeline": read_timeline(workbook["Correlated Timeline"], known),
        "gfma_state_history": read_state_history(
            workbook["GFM-A State History"], section.group(1) if section else "GFM-A", known
        ),
        "method": read_method(workbook["Method & Tools"]),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", type=Path, help="Analysis scenarios workbook (.xlsx)")
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT, help="JSON file to write")
    args = parser.parse_args()

    document = extract(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        f"Wrote {args.output}: {len(document['scenarios'])} scenarios, {len(document['test_cases'])} test cases, "
        f"{len(document['timeline'])} timeline events, {len(document['gfma_state_history']['states'])} GFM-A states, "
        f"{len(document['method']['steps'])} method steps, {len(document['method']['open_questions'])} open questions"
    )


if __name__ == "__main__":
    main()
